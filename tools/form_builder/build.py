import os, shutil
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

SRC = '/mnt/user-data/uploads/교정본부_자료요청_통합서식_자료요청_서식.xlsx'
OUT = '/home/claude/out/교정본부_자료요청_취합본_2026-10-21.xlsx'
R = '/home/claude/received/'
os.makedirs('/home/claude/out', exist_ok=True)
shutil.copy(SRC, OUT)

def rf(prefix):
    return R + [f for f in os.listdir(R) if f.startswith(prefix)][0]

wb = load_workbook(OUT)
BLUE = Font(name='Calibri', size=11, color='FF0000FF')
BLUE_T = Font(name='맑은 고딕', size=9, color='FF0000FF')
RED_T = Font(name='맑은 고딕', size=9, color='FFC00000')
F = '맑은 고딕'
thin = Side(style='thin', color='FFBFBFBF')
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
HDR = PatternFill('solid', fgColor='FF1F3864')
SEC = PatternFill('solid', fgColor='FFD9E1F2')
YEL = PatternFill('solid', fgColor='FFFFFF00')
GRY = PatternFill('solid', fgColor='FFEDEDED')

def put(ws, ref, v, font=BLUE, fmt=None):
    c = ws[ref]; c.value = v; c.font = font
    if fmt: c.number_format = fmt
    return c

# ---------- 1.요청표: 소관(회신 부서)·진행 상태 ----------
ws = wb['1.요청표']
owner = {5: '교정본부 복지과', 6: '교정본부 보안과', 7: '교정본부 보안과', 8: '마약사범재활팀', 9: '마약사범재활팀',
         10: '교정본부 복지과', 11: '교정본부 의료과', 12: '마약사범재활팀', 13: '마약사범재활팀', 14: '마약사범재활팀',
         15: '각 소관과', 16: '교정기획과', 17: '복지부·범죄예방정책국·치료처우과'}
status = {5: '부분 회신 (10항목 중 8)', 6: '부분 회신 (연인원 대신 연말현원)', 7: '거의 미회신 (연말현원만)',
          8: '부분 회신 (이수 인원만)', 9: '회신 완료', 10: '총액만 회신 (비목 산출 불가)',
          11: '부분 회신 (단위 확인 필요)', 12: '미회신', 13: '문헌 대체 (김도균·나종민 2025)',
          14: '미회신', 15: '미회신', 16: '미회신 (서식 시트 없음)', 17: '부분 회신 (복지부 미회신)'}
for r in owner:
    ws.cell(r, 5).value = owner[r]
    c = ws.cell(r, 7); c.value = status[r]
    c.font = Font(name=F, size=9, bold=True, color='FFC00000' if ('미회신' in status[r] or '거의' in status[r]) else ('FF00632E' if '완료' in status[r] else 'FF9C5700'))
ws['A2'].value = ('총 13건(교정본부 12 + 타 기관 1). 2026-10-21 취합 기준: 완료 1, 부분 6, 미회신 5, 문헌 대체 1(서식9). '
                  '상세는 「취합현황」 시트, 계산은 「원단위계산」 시트.')

# ---------- 서식1 ----------
s = load_workbook(rf('b6e46f2f'))['서식1 세출예산']
ws = wb['서식1 세출예산']
for r in range(7, 17):
    v = s.cell(r, 4).value
    if v is not None:
        put(ws, f'D{r}', v, fmt='#,##0'); put(ws, f'E{r}', '회신(복지과, 26.9.7)', BLUE_T)
    else:
        put(ws, f'E{r}', '미회신', RED_T)
put(ws, 'A20', '회신 없음: 8번(교정시설 취득)·10번(교도작업특별회계). 별도 회신(hwpx): 교도소 일반회계 세출 결산 2024년 1,674,128,000,000원, 2025년 1,699,217,000,000원.', BLUE_T)

# ---------- 서식2 ----------
s = load_workbook(rf('7c181003'))['서식2 수용인원']
ws = wb['서식2 수용인원']
for r, v in [(7, s['D7'].value), (9, s['D9'].value), (10, s['D10'].value), (11, s['D11'].value), (12, s['D12'].value)]:
    put(ws, f'D{r}', v, fmt='#,##0.0' if r == 10 else '#,##0'); put(ws, f'E{r}', '회신(보안과)', BLUE_T)
put(ws, 'E8', '미회신: 회신값 64,626명은 25.12.31 연말 현원(연인원 아님)', RED_T)
put(ws, 'B19', '기관별 54개 기관 값은 「서식2-기관별(회신)」 시트로 옮김', BLUE_T)
inst = []
for r in range(79, 133):
    if s.cell(r, 2).value: inst.append((s.cell(r, 2).value, s.cell(r, 4).value))
ws2 = wb.create_sheet('서식2-기관별(회신)', wb.sheetnames.index('서식2 수용인원') + 1)
ws2['A1'] = '서식 2 [추가] 기관별 1일 평균 수용인원 (2025년, 보안과 회신)'; ws2['A1'].font = Font(name=F, size=12, bold=True)
for j, h in enumerate(['번호', '기관명', '1일 평균 수용인원(명)', '비고 (모델 KEY)'], 1):
    c = ws2.cell(3, j, h); c.font = Font(name=F, size=9, bold=True, color='FFFFFFFF'); c.fill = HDR; c.border = BOX
for i, (n, v) in enumerate(inst, 1):
    ws2.cell(3 + i, 1, i).font = Font(name=F, size=9)
    ws2.cell(3 + i, 2, n).font = Font(name=F, size=9)
    c = ws2.cell(3 + i, 3, v); c.font = BLUE; c.number_format = '#,##0.0'
    ws2.cell(3 + i, 4, 'IN-N-02  기관별 분모').font = Font(name=F, size=9, color='FF7F7F7F')
    for j in range(1, 5): ws2.cell(3 + i, j).border = BOX
tr = 4 + len(inst)
ws2.cell(tr, 2, '합계').font = Font(name=F, size=9, bold=True)
ws2.cell(tr, 3, f'=SUM(C4:C{tr-1})').number_format = '#,##0'
ws2.cell(tr + 1, 2, '전국 1일 평균(서식2)').font = Font(name=F, size=9)
ws2.cell(tr + 1, 3, "='서식2 수용인원'!D7").number_format = '#,##0'
ws2.cell(tr + 1, 3).font = Font(name='Calibri', size=11, color='FF008000')
ws2.cell(tr + 2, 2, '차이(합계−전국)').font = Font(name=F, size=9)
ws2.cell(tr + 2, 3, f'=C{tr}-C{tr+1}').number_format = '#,##0;(#,##0);-'
for col, w in zip('ABCD', [6, 16, 20, 22]): ws2.column_dimensions[col].width = w

# ---------- 서식3 ----------
ws = wb['서식3 마약류사범']
put(ws, 'E7', '미회신: 1일 평균 없음. 25.12.31 연말 현원 7,429명만 회신', RED_T)
put(ws, 'E9', '미회신(회신 시트에서 이 칸에 연말 현원 7,429명이 기재됨)', RED_T)
for r in [8, 10, 11, 12, 13, 14, 15]: put(ws, f'E{r}', '미회신', RED_T)
put(ws, 'B19', '사유(보안과): 해당 통계는 공식 통계로 집계 및 제공하고 있지 않음', BLUE_T)

# ---------- 서식4 ----------
s = load_workbook(rf('1de64b66'))['서식4 재활교육 실적']
ws = wb['서식4 재활교육 실적']
for r in range(8, 20):
    put(ws, f'B{r}', '전국 합계', BLUE_T)
    put(ws, f'G{r}', s[f'G{r}'].value, fmt='#,##0')
    if s[f'A{r}'].value == '2026. 7.': put(ws, f'A{r}', '2026(7월까지)', BLUE_T)
put(ws, 'A32', '회신(마약사범재활팀): 이수 인원만 전국 합계로 회신. 기수·개시·중도탈락은 "별도 집계 관리하고 있지 않은 항목에는 미응답". 같은 내용 파일 3개 수령.', BLUE_T)

# ---------- 서식5 ----------
s = load_workbook(rf('c6f72bc7'))['서식5 전담인력']
ws = wb['서식5 전담인력']
for r in range(8, 12):
    for col in 'BDEFG':
        v = s[f'{col}{r}'].value
        put(ws, f'{col}{r}', v, BLUE if col in 'EFG' else BLUE_T)
    put(ws, f'I{r}', s[f'I{r}'].value, BLUE_T)
put(ws, 'I7', '직급별 구성(회신 비고)', BLUE_T)

# ---------- 서식6 ----------
s = load_workbook(rf('b6e46f2f'))['서식6 재활예산']
ws = wb['서식6 재활예산']
for r in range(7, 15): put(ws, f'E{r}', '산출 불가', RED_T)
put(ws, 'A15', '합계', Font(name=F, size=9, bold=True))
put(ws, 'B15', '마약사범 재활 관련 예산 총액 (비목 구분 불가)', BLUE_T)
put(ws, 'C15', '원', Font(name=F, size=9))
put(ws, 'D15', 5999392000, fmt='#,##0')
put(ws, 'E15', '회신(복지과): 총액만', BLUE_T)
put(ws, 'B18', s['B18'].value, BLUE_T)
put(ws, 'B19', s['B19'].value + '  → 예산액인지 집행액인지, 시설(설치)비 포함 여부 확인 필요', BLUE_T)

# ---------- 서식7 ----------
s = load_workbook(rf('d1e168c6'))['서식7 의료·검사']
ws = wb['서식7 의료·검사']
for r in [7, 8, 16, 17]:
    for col in 'DEFGH':
        put(ws, f'{col}{r}', s[f'{col}{r}'].value, fmt='#,##0')
for r in [7, 8]: put(ws, f'I{r}', '단위 확인 필요: "원"으로 표기됐으나 천원으로 추정', RED_T)
put(ws, 'B16', '연간 진료 건수(내부)', Font(name=F, size=9))
for r in [9, 10, 11, 12, 13, 14, 15, 18, 19]: put(ws, f'I{r}', '미생산 통계(의료과)', RED_T)

# ---------- 별도요청 ----------
ws = wb['별도요청 타기관']
s1 = load_workbook(rf('7c181003'))['별도요청 타기관']
s2 = load_workbook(rf('6b29e3a2'))['별도요청 타기관']
for r in [11, 12]:
    for col in 'DEFGH': put(ws, f'{col}{r}', s1[f'{col}{r}'].value, fmt='#,##0')
    put(ws, f'I{r}', '회신(범죄예방정책국)', BLUE_T)
put(ws, 'I13', '산출 불가(범죄예방정책국)', RED_T)
for col in 'DEFGH':
    v = s2[f'{col}14'].value
    put(ws, f'{col}14', int(v.replace(',', '').replace('백만원', '')) * 1_000_000, fmt='#,##0')
    put(ws, f'{col}15', s2[f'{col}15'].value, fmt='#,##0')
put(ws, 'I14', '회신(치료처우과): 병원 전체 결산', BLUE_T)
put(ws, 'I15', '회신(치료처우과)', BLUE_T)
for r in [7, 8, 9, 10]: put(ws, f'I{r}', '미회신(복지부)', RED_T)
put(ws, 'B21', '범죄예방정책국: ' + s1['B21'].value, BLUE_T)
put(ws, 'B22', '복지부 미회신. 치료보호 인원은 2025 마약류 범죄백서 표 3-82로 대체 가능(2021~2025: 280, 421, 641, 875, 1,649명). 예산은 대체 자료 없음.', BLUE_T)

wb.save(OUT)
print('saved', OUT, len(inst))
