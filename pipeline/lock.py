"""서식 잠금: 엑셀의 고정 칸(제목·항목명·산식·KEY 등)과 잠금 파일 해시를 만들고 점검한다."""
import hashlib, re
from openpyxl import load_workbook
from common import (ROOT, FORM_LOCK, MANIFEST, LOCKED_FILES, CALC_SHEET, GENERATED_SHEETS, GENERATED_CELLS,
                    FULL_LOCK_SHEETS, INPUT_FONT, INPUT_FILL, load_json)


def _rgb(c):
    try:
        v = c.font.color.rgb if c.font and c.font.color is not None else None
        return v if isinstance(v, str) else None
    except Exception:
        return None


def _fill(c):
    try:
        v = c.fill.fgColor.rgb if c.fill and c.fill.fill_type == 'solid' else None
        return v if isinstance(v, str) else None
    except Exception:
        return None


def norm(v):
    """비교용 표기. 산식은 대소문자·공백·$를 무시(엑셀 저장 시 표기만 바뀌는 경우 대비)."""
    if v is None:
        return ''
    if isinstance(v, str):
        s = v.strip()
        if s.startswith('='):
            return '=' + re.sub(r'[\s$]', '', s[1:]).upper()
        return s
    if isinstance(v, bool):
        return str(v)
    if isinstance(v, (int, float)):
        return repr(round(float(v), 9))
    return str(v)


def is_input_cell(sheet, c):
    """사람이 값을 넣어도 되는 칸인지 판정(잠금 기준을 만들 때만 사용)."""
    if sheet in FULL_LOCK_SHEETS:
        return False
    if _rgb(c) in INPUT_FONT or _fill(c) in INPUT_FILL:
        return not (isinstance(c.value, str) and c.value.startswith('='))
    if sheet == CALC_SHEET:
        return False  # 계산 시트는 파란 글씨·노란 칸만 입력 가능
    return c.value is None  # 서식 시트의 빈칸은 회신값 입력용


def build_form_lock(path):
    wb = load_workbook(path)
    out = {'sheets': wb.sheetnames, 'locked': {}, 'extent': {}}
    for ws in wb.worksheets:
        if ws.title in GENERATED_SHEETS:
            continue
        gen = set(GENERATED_CELLS.get(ws.title, []))
        cells = {}
        for row in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=ws.max_column):
            for c in row:
                if c.coordinate in gen or type(c).__name__ == 'MergedCell':
                    continue
                if not is_input_cell(ws.title, c):
                    cells[c.coordinate] = norm(c.value)
        out['locked'][ws.title] = cells
        if ws.title in FULL_LOCK_SHEETS or ws.title == CALC_SHEET:
            out['extent'][ws.title] = [ws.max_row, ws.max_column]
    return out


def check_form(path):
    """현재 엑셀이 잠금 기준과 맞는지. 어긋난 칸 목록을 돌려준다."""
    lock = load_json(FORM_LOCK)
    wb = load_workbook(path)
    errs = []
    if wb.sheetnames != lock['sheets']:
        errs.append(f'시트 구성이 바뀌었습니다. 기준: {lock["sheets"]} / 현재: {wb.sheetnames}')
    for name, cells in lock['locked'].items():
        if name not in wb.sheetnames:
            errs.append(f'[{name}] 시트가 없습니다')
            continue
        ws = wb[name]
        bad = [(k, v, norm(ws[k].value)) for k, v in cells.items() if norm(ws[k].value) != v]
        for k, v, now in bad[:15]:
            errs.append(f'[{name}] {k} 고정 칸이 바뀌었습니다. 기준 "{v[:40]}" → 현재 "{now[:40]}"')
        if len(bad) > 15:
            errs.append(f'[{name}] 외 {len(bad) - 15}칸 더')
        if name in lock['extent']:
            mr, mc = lock['extent'][name]
            extra = [c.coordinate for row in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=ws.max_column) for c in row
                     if c.value is not None and (c.row > mr or c.column > mc)]
            if extra:
                errs.append(f'[{name}] 잠긴 영역 밖에 새 내용이 있습니다: {", ".join(extra[:10])}')
    return errs


def file_hashes():
    return {f: hashlib.sha256((ROOT / f).read_bytes()).hexdigest() for f in LOCKED_FILES}


def check_manifest():
    m = load_json(MANIFEST)
    now = file_hashes()
    errs = []
    for f, h in now.items():
        if m['files'].get(f) != h:
            errs.append(f'잠금 파일이 바뀌었습니다: {f}')
    for f in m['files']:
        if f not in now:
            errs.append(f'잠금 목록의 파일이 없습니다: {f}')
    return errs, m.get('form_version')
