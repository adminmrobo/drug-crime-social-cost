from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.comments import Comment

OUT = '/home/claude/out/교정본부_자료요청_취합본_2026-10-21.xlsx'
wb = load_workbook(OUT)
F = '맑은 고딕'
thin = Side(style='thin', color='FFBFBFBF'); BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
HDR = PatternFill('solid', fgColor='FF1F3864'); SEC = PatternFill('solid', fgColor='FFD9E1F2')
YEL = PatternFill('solid', fgColor='FFFFFF00'); KEYF = PatternFill('solid', fgColor='FFFCE4D6')
COL = {'input': 'FF0000FF', 'link': 'FF008000', 'formula': 'FF000000', 'assump': 'FF0000FF'}

ws = wb.create_sheet('원단위계산', wb.sheetnames.index('1.요청표') + 1)
ws.sheet_view.showGridLines = False
for col, w in zip('ABCDEFGHIJ', [12, 44, 18, 12, 62, 14, 14, 14, 14, 14]):
    ws.column_dimensions[col].width = w
ws['A1'] = '원단위·증분비용·손익분기 계산 (2025년 기준, 2026-10-21 취합본)'
ws['A1'].font = Font(name=F, size=13, bold=True)
ws['A2'] = ('글자색: 파랑=직접 입력값(회신·문헌), 초록=다른 시트 연결, 검정=계산식. 노란 칸=연구진 가정(바꾸면 전체 재계산). '
            '주황 칸=핵심 결과. UC-10(치료개입)·UC-12(사후관리)는 자료 미비로 0 처리 → ΔC는 하한값.')
ws['A2'].font = Font(name=F, size=9, color='FF595959'); ws['A2'].alignment = Alignment(wrap_text=True)
ws.merge_cells('A2:E2'); ws.row_dimensions[2].height = 30

N = {}; r = [3]

def sec(title):
    r[0] += 1
    for c in range(1, 6):
        ws.cell(r[0], c).fill = SEC
    ws.cell(r[0], 1, title).font = Font(name=F, size=10, bold=True)
    r[0] += 1

def ref(s):
    for k in sorted(N, key=len, reverse=True):
        s = s.replace('{' + k + '}', N[k])
    return s

def row(key, label, val, unit, note, kind='formula', fmt='#,##0', name=None, hl=False):
    rr = r[0]
    ws.cell(rr, 1, key).font = Font(name=F, size=9, color='FF7F7F7F')
    ws.cell(rr, 2, label).font = Font(name=F, size=9, bold=hl)
    v = ref(val.replace('{{ROW}}', str(rr))) if isinstance(val, str) else val
    c = ws.cell(rr, 3, v); c.number_format = fmt
    c.font = Font(name='Calibri', size=11, color=COL[kind], bold=hl)
    if kind == 'assump': c.fill = YEL
    if hl: c.fill = KEYF
    ws.cell(rr, 4, unit).font = Font(name=F, size=9)
    n = ws.cell(rr, 5, note); n.font = Font(name=F, size=8, color='FF595959'); n.alignment = Alignment(wrap_text=True, vertical='top')
    for cc in range(1, 6): ws.cell(rr, cc).border = BOX
    if name: N[name] = f'$C${rr}'
    r[0] += 1
    return rr

S1, S2, S4, S5, S6, S7, SX = ("'서식1 세출예산'", "'서식2 수용인원'", "'서식4 재활교육 실적'", "'서식5 전담인력'",
                              "'서식6 재활예산'", "'서식7 의료·검사'", "'별도요청 타기관'")

# ---------------- A. 투옥 ----------------
sec('A. 투옥 원가중심점 (IN-C / IN-N) → UC-01 총원가, UC-02 변동원가')
row('IN-N-01', '1일 평균 수용인원(전국, 2025)', f'={S2}!D7', '명', '서식2 1번(보안과)', 'link', name='pop')
row('', '연 일수', 365, '일', '가정', 'assump', name='days')
row('IN-C-02', '급식비', f'={S1}!D7', '원', '서식1 1번(복지과)', 'link', name='food')
row('IN-C-03', '피복·침구비', f'={S1}!D8', '원', '서식1 2번', 'link', name='cloth')
row('IN-C-04', '공공요금', f'={S1}!D9', '원', '서식1 3번', 'link', name='util')
row('IN-C-05', '일반수용비·기타 물건비', f'={S1}!D10', '원', '서식1 4번', 'link', name='gen')
row('IN-C-06', '교도소운영 인건비', f'={S1}!D11', '원', '서식1 5번', 'link', name='labor')
row('PG-C-01', '교정·교화 사업', f'={S1}!D12', '원', '서식1 6번. 재활예산(서식6)을 포함하므로 UC-01에 넣으면 재활비가 이중으로 잡힘 → 아래 UC-01a에서는 제외', 'link', name='edu')
row('IN-C-08', '교정시설장비 운용·현대화', f'={S1}!D13', '원', '서식1 7번', 'link', name='equip')
row('IN-C-10', '교정본부·지방교정청 공통관리비', f'={S1}!D15', '원', '서식1 9번', 'link', name='hq')
row('', '교도소 일반회계 세출 결산 총액(2025)', 1699217000000, '원', '복지과 별도 회신(hwpx "자료요청 2차"). 의료비·직업훈련 등 서식1 미분류 항목 포함', 'input', name='total')
row('', '회신 항목 합계(교정·교화 제외)', '={food}+{cloth}+{util}+{gen}+{labor}+{equip}+{hq}', '원', '교정시설 취득(8번)·교도작업특별회계(10번) 미회신으로 빠짐', name='sumA')
row('UC-01a', '수용 1인·연 총원가 (회신 항목 기준)', '={sumA}/{pop}', '원/인·년', '하한. 보고서 초안 2,396만 원은 교정·교화 포함값', name='uc01y', hl=True)
row('UC-01a', '수용 1인·1일 총원가 (회신 항목 기준)', '={uc01y}/{days}', '원/인·일', '', name='uc01d')
row('UC-01b', '수용 1인·연 총원가 (일반회계 결산 총액 기준)', '=({total}-{edu})/{pop}', '원/인·년', '상한. 결산 총액에서 교정·교화(재활 포함)만 제외', name='uc01by', hl=True)
row('', '변동원가에 넣을 공공요금 비율', 0.5, '', '가정: 공공요금은 수용인원에 반쯤 비례(준변동비)', 'assump', fmt='0%', name='utilshare')
row('UC-02', '수용 1인·1일 변동원가(한계원가)', '=({food}+{cloth}+{gen}+{util}*{utilshare})/{pop}/{days}', '원/인·일', '급식+피복+일반수용비+공공요금×비율. 인건비·장비·관리비는 단기 고정비로 봄', name='uc02', hl=True)

# ---------------- B. 프로그램 ----------------
sec('B. 투입 원가중심점 (PG-C / PG-N / PG-D) → UC-07 인시당 원가, UC-06 과정별 1인당 원가')
row('PG-C-02·03', '마약사범 재활 예산 총액(2025)', f'={S6}!D15', '원', '서식6(복지과). 비목 구분 불가, 예산/집행 구분과 시설비 포함 여부 확인 필요', 'link', name='pgbud')
# personnel table
pr = r[0]
heads = ['직급', '화성직훈(교)', '부산(교)', '청주여자(교)', '광주(교)', '월 봉급(원)']
for j, h in enumerate(heads):
    c = ws.cell(pr, 2 + j, h); c.font = Font(name=F, size=9, bold=True, color='FFFFFFFF'); c.fill = HDR; c.border = BOX
ws.cell(pr, 1, 'PG-C-04').font = Font(name=F, size=9, color='FF7F7F7F')
grades = [('5급', [1, 1, 1, 1], 4066800), ('6급', [3, 2, 1, 2], 3468700), ('7급', [1, 4, 1, 3], 3133300),
          ('8급', [2, 4, 3, 2], 2813000), ('9급', [4, 0, 5, 3], 2542700)]
for i, (g, cnt, pay) in enumerate(grades, 1):
    rr = pr + i
    ws.cell(rr, 2, g).font = Font(name=F, size=9)
    for j, v in enumerate(cnt):
        c = ws.cell(rr, 3 + j, v); c.font = Font(name='Calibri', size=11, color=COL['input'])
    c = ws.cell(rr, 7, pay); c.font = Font(name='Calibri', size=11, color=COL['input']); c.number_format = '#,##0'
    for cc in range(2, 8): ws.cell(rr, cc).border = BOX
rr = pr + 6
ws.cell(rr, 2, '인원 합계').font = Font(name=F, size=9, bold=True)
for j in range(4):
    col = 'CDEF'[j]; ws.cell(rr, 3 + j, f'=SUM({col}{pr+1}:{col}{pr+5})').font = Font(name='Calibri', size=11)
rr = pr + 7
ws.cell(rr, 2, '프로그램 직접투입 비율').font = Font(name=F, size=9)
for j, (col, src) in enumerate(zip('CDEF', ['G8', 'G9', 'G10', 'G11'])):
    c = ws.cell(rr, 3 + j, f'={S5}!{src}/100'); c.number_format = '0.0%'; c.font = Font(name='Calibri', size=11, color=COL['link'])
ws.cell(pr + 8, 2, '연 인건비(봉급×12×계수)').font = Font(name=F, size=9)
ws.cell(pr + 9, 2, '직접투입 인건비').font = Font(name=F, size=9)
r[0] = pr + 10
row('', '봉급 외 인건비 계수(수당·상여·복리·연금부담 포함)', 1.6, '배', '가정: 연간 총인건비 ≈ 연 봉급×1.6. 보고서 초안의 24.3억 원은 계수 약 1.5에 해당', 'assump', fmt='0.00', name='pcoef')
for j, col in enumerate('CDEF'):
    c = ws.cell(pr + 8, 3 + j, f'=SUMPRODUCT({col}{pr+1}:{col}{pr+5},$G${pr+1}:$G${pr+5})*12*{N["pcoef"]}')
    c.number_format = '#,##0'; c.font = Font(name='Calibri', size=11)
    c = ws.cell(pr + 9, 3 + j, f'={col}{pr+8}*{col}{pr+7}'); c.number_format = '#,##0'; c.font = Font(name='Calibri', size=11)
for rr2 in range(pr + 6, pr + 10):
    for cc in range(2, 7): ws.cell(rr2, cc).border = BOX
ws.cell(pr, 8, '출처: 직급 구성=서식5 비고(마약사범재활팀, 2026 정원). 월 봉급=인사혁신처 2026 일반직 봉급표 10호봉(가정).').font = Font(name=F, size=8, color='FF595959')
ws.cell(pr, 8).alignment = Alignment(wrap_text=True, vertical='top'); ws.merge_cells(start_row=pr, start_column=8, end_row=pr + 4, end_column=10)
row('PG-C-04', '전담인력 연 인건비 합계(44명)', f'=SUM(C{pr+8}:F{pr+8})', '원', '2026년 4개 마약사범재활과 정원 기준', name='ptot')
row('PG-D-02', '그중 프로그램 직접투입분', f'=SUM(C{pr+9}:F{pr+9})', '원', '기관별 투입비율(64.0~68.2%) 적용', name='pdir')
row('', '2025년 원가에 전담인력 인건비 반영(1=예, 0=아니오)', 1, '', '가정: 4개 과는 2026.1 신설이므로 2025년에는 다른 인력이 수행. 2026 구조를 2025 대리값으로 씀', 'assump', fmt='0', name='pswitch')
row('', 'PG 비용풀(재활예산 + 직접투입 인건비)', '={pgbud}+{pdir}*{pswitch}', '원', '', name='pgpool', hl=True)

# course table
cr = r[0] + 1
heads = ['과정', '이수 인원(2025)', '과정 시간', '이수 인시', 'UC-06 1인당 원가', 'ΔC 1인당 증분비용', '손익분기 SCᵣ*', '부분 B/C']
for j, h in enumerate(heads):
    c = ws.cell(cr, 2 + j, h); c.font = Font(name=F, size=9, bold=True, color='FFFFFFFF'); c.fill = HDR; c.border = BOX
    c.alignment = Alignment(wrap_text=True, horizontal='center')
ws.row_dimensions[cr].height = 28
courses = [('기본', 12), ('집중', 13), ('심화', 14), ('회복이음', 15)]
for i, (nm, sr) in enumerate(courses, 1):
    rr = cr + i
    ws.cell(rr, 1, 'PG-N-02').font = Font(name=F, size=9, color='FF7F7F7F')
    ws.cell(rr, 2, nm).font = Font(name=F, size=9)
    c = ws.cell(rr, 3, f'={S4}!G{sr}'); c.font = Font(name='Calibri', size=11, color=COL['link']); c.number_format = '#,##0'
    c = ws.cell(rr, 4, f'={S4}!D{sr}'); c.font = Font(name='Calibri', size=11, color=COL['link'])
    ws.cell(rr, 5, f'=C{rr}*D{rr}').number_format = '#,##0'
    for cc in range(2, 10): ws.cell(rr, cc).border = BOX
rr = cr + 5
ws.cell(rr, 2, '합계·평균').font = Font(name=F, size=9, bold=True)
ws.cell(rr, 3, f'=SUM(C{cr+1}:C{cr+4})').number_format = '#,##0'
ws.cell(rr, 5, f'=SUM(E{cr+1}:E{cr+4})').number_format = '#,##0'
for cc in range(2, 10): ws.cell(rr, cc).border = BOX
course_rows = (cr + 1, cr + 4, rr)
r[0] = rr + 2
N['hours'] = f'$E${rr}'; N['comp'] = f'$C${rr}'
row('UC-07', '재활교육 1인시당 원가', '={pgpool}/{hours}', '원/인시', 'PG 비용풀 ÷ 이수 인시. 개시·중도탈락 인원 미회신으로 이수 인원만 분모 → 과대 추정 쪽', name='uc07', hl=True)
row('UC-06', '이수자 1인당 평균 프로그램 원가', '={pgpool}/{comp}', '원/인', '과정 구성 가중평균. 보고서 초안 377.6만 원(예산만)·530.5만 원(인건비 포함)과 대조', name='uc06avg', hl=True)

# ---------------- C. 치료·사후관리 ----------------
sec('C. 치료(UC-10)·사후관리(UC-12) 원가중심점 — 자료 미비, 참고값만 산출')
row('', '서식7 의료비 단위 환산 계수', 1000, '배', '가정: 회신값이 천원 단위로 보임(원 단위면 전국 의료비 342만 원이 되어 불가능). 의료과 확인 후 수정', 'assump', fmt='#,##0', name='medunit')
row('TX-C-01', '수용자 의료비 총액(2025)', f'={S7}!H7*{{medunit}}', '원', '서식7 1번(의료과)', 'link', name='med')
row('참고', '수용자 1인당 연 의료비(전체 수용자)', '={med}/{pop}', '원/인·년', '마약류사범 치료개입 원가(UC-10)가 아님. 치료개입 대상 인원·정신과 약제비 미생산', name='medpc')
row('UC-10', 'ΔC에 넣는 치료개입 원가', 0, '원/인', '자동: 서식7의 정신과 진료·약품비(6번)와 치료개입 대상 인원(12번)이 모두 숫자로 들어오면 그 몫으로 계산. 없으면 F열(수기값, 기본 0)', name='uc10')
row('CU-C-01', '국립법무병원 결산(2025)', f'={SX}!H14', '원', '치료처우과 회신. 병원 전체 결산이라 약물중독(2호) 몫을 나눌 분모(병원 전체 연평균 수용인원)가 필요', 'link', name='nfh')
row('CU-N-02', '보호관찰 대상 마약류사범(2025)', f'={SX}!H11', '명', '범죄예방정책국 회신. 관련 예산은 "산출 불가"', 'link', name='prob')
row('UC-12', 'ΔC에 넣는 사후관리 원가', 0, '원/인', '자동: 별도요청 2번(치료보호 집행액)·3번(지원 인원) 2025년 값이 들어오면 1인·연 원가×연계비율(G열)÷(1+r). 없으면 F열(수기값, 기본 0)', name='uc12')
row('ΔM', '재활 참여로 달라지는 수용기간', 0, '일', '가정: 차이 없음(치료조건부 가석방 도입 시 음수)', 'assump', fmt='#,##0', name='dm')

# ---------------- D. 효과 ----------------
sec('D. 효과 (OC) — 서식9 회신 전에는 김도균·나종민(2025) 표본(F열)을 씀')
row('OC-03', '서식9 회신값 사용 여부(1=서식9, 0=논문)', "=IF(AND(N('서식9 이수별 재복역'!D9)>0,N('서식9 이수별 재복역'!E9)>0,N('서식9 이수별 재복역'!D10)>0,N('서식9 이수별 재복역'!E10)>0),1,0)", '', '자동 판정', fmt='0', name='oc9')
row('OC-03', '이수 집단 인원', "=IF({oc9}=1,'서식9 이수별 재복역'!D9,F{{ROW}})", '명', '자동: 서식9의 2025 전체 이수·미이수 두 줄이 모두 채워지면 서식9 값, 아니면 F열(김도균·나종민 2025 표5)', name='n1')
row('OC-03', '이수 집단 재복역', "=IF({oc9}=1,'서식9 이수별 재복역'!E9,F{{ROW}})", '명', 'n1과 같은 규칙', name='e1')
row('OC-03', '미이수 집단 인원', "=IF({oc9}=1,'서식9 이수별 재복역'!D10,F{{ROW}})", '명', 'n1과 같은 규칙', name='n0')
row('OC-03', '미이수 집단 재복역', "=IF({oc9}=1,'서식9 이수별 재복역'!E10,F{{ROW}})", '명', 'n1과 같은 규칙', name='e0')
row('r0', '기저재범률(미이수 재복역률)', '={e0}/{n0}', '', '추적 약 5~6.7년 누적', fmt='0.0%', name='r0', hl=True)
row('r1', '이수 집단 재복역률', '={e1}/{n1}', '', '', fmt='0.0%', name='r1')
row('Δr', '재범감소폭(상대 감소율) = (r0−r1)/r0', '=({r0}-{r1})/{r0}', '', '단순 비교값. 논문 Cox 모형 위험비 0.805(통제변수 반영, p<.05) → 약 19.5%로 비슷. 로지스틱(오즈비 0.768)은 p=.066으로 유의하지 않음', fmt='0.0%', name='dr', hl=True)
row('', '1인당 회피 재범 확률 = r0×Δr', '={r0}*{dr}', '', 'r0 − r1 과 같음', fmt='0.0%', name='avoid')

# ---------------- E. 할인 ----------------
sec('E. 재범 시점 분포 hₜ와 할인계수 D (r = 4.5%, T = 5년)')
row('', '사회적 할인율 r', 0.045, '', '인계 메모 기준', 'assump', fmt='0.0%', name='rate')
row('OC-05', '전체 재복역 중 3년 안에 일어나는 비중', 0.363 / 0.456, '', '가정: 2018 출소 마약류사범 3년 내 재복역률 36.3%(법무부, 논문 재인용) ÷ 표본 장기 재복역률 45.6%. 서식8 경과기간 자료 오면 교체', 'assump', fmt='0.0%', name='s3')
hr = r[0]
for t in range(1, 6):
    hv = '={s3}/3' if t <= 3 else '=(1-{s3})/2'
    row('hₜ', f'{t}년차 재범 비중 h{t}', hv, '', '1~3년 균등, 4~5년 균등 배분(가정)', fmt='0.0%', name=f'h{t}')
row('D', '할인계수 D = Σ hₜ/(1+r)^t', '={h1}/(1+{rate})^1+{h2}/(1+{rate})^2+{h3}/(1+{rate})^3+{h4}/(1+{rate})^4+{h5}/(1+{rate})^5', '', '', fmt='0.000', name='D', hl=True)

# ---------------- F. 부분 SCr ----------------
sec('F. 재범 1건당 사회적 비용 SCᵣ — 이번에 확보된 형사사법 부분만 (하한)')
sr = r[0]
for lab, cnt, mid in [('1년 미만', 290, 0.5), ('1~2년', 837, 1.5), ('2~3년', 444, 2.5), ('3~5년', 318, 4), ('5년 이상', 237, 6)]:
    row('', f'향정 실형 {lab}: 인원 / 구간 대표값(년)', cnt, '명', '2025 마약류 범죄백서 표 3-69', 'input', name=f'n_{lab}')
    ws.cell(r[0] - 1, 6, mid).font = Font(name='Calibri', size=11, color=COL['assump']); ws.cell(r[0] - 1, 6).fill = YEL
    ws.cell(r[0] - 1, 6).number_format = '0.0'
ws.cell(sr - 1, 6, '대표값(년)').font = Font(name=F, size=9, bold=True)
row('', '평균 실형 기간(선고 기준)', f'=SUMPRODUCT(C{sr}:C{sr+4},F{sr}:F{sr+4})/SUM(C{sr}:C{sr+4})', '년', '구간 대표값은 가정(노란 칸). 가석방으로 실제 복역기간은 이보다 짧음', fmt='0.00', name='sent')
row('', '실제 복역 비율(선고 대비)', 0.85, '', '가정: 가석방 등 반영. 교정통계연보 자료로 교체 필요', 'assump', fmt='0%', name='served')
row('', '재수용 비용 = 복역기간 × UC-01a', '={sent}*{served}*{uc01y}', '원/건', '재복역 1건이 다시 일으키는 수용비. 복역 기간 중 할인은 생략(약간 과대)', name='reinc')
row('', '검찰 세출 결산(2024)', 1216261000000, '원', '대검찰청 정보공개 회신(hwpx)', 'input', name='prosb')
row('', '형사사건 검거인원(2024)', 1405185, '명', '같은 회신', 'input', name='arrest')
row('', '검찰 단계 1인당 비용', '={prosb}/{arrest}', '원/인', '전체 형사사건 평균. 경찰·법원 비용은 아직 없음', name='prospc')
row('SCᵣ(부분)', '재범 1건당 형사사법 비용(재수용+검찰)', '={reinc}+{prospc}', '원/건', '의료·생산성·사회혼란·고통비용은 3장 추계 후 추가 → 지금 값은 하한', name='scr', hl=True)

# ---- course table formulas ----
for rr in range(course_rows[0], course_rows[1] + 1):
    ws.cell(rr, 6, ref(f'=D{rr}*{{uc07}}')).number_format = '#,##0'
    ws.cell(rr, 7, ref(f'=F{rr}+{{uc10}}+{{uc12}}+{{dm}}*{{uc02}}')).number_format = '#,##0'
    ws.cell(rr, 8, ref(f'=G{rr}/({{avoid}}*{{D}})')).number_format = '#,##0'
    ws.cell(rr, 9, ref(f'={{avoid}}*{{scr}}*{{D}}/G{rr}')).number_format = '0.00'
rr = course_rows[2]
ws.cell(rr, 6, ref('={uc06avg}')).number_format = '#,##0'
ws.cell(rr, 7, ref(f'=F{rr}+{{uc10}}+{{uc12}}+{{dm}}*{{uc02}}')).number_format = '#,##0'
ws.cell(rr, 8, ref(f'=G{rr}/({{avoid}}*{{D}})')).number_format = '#,##0'
ws.cell(rr, 9, ref(f'={{avoid}}*{{scr}}*{{D}}/G{rr}')).number_format = '0.00'
for c in range(6, 10):
    ws.cell(rr, c).fill = KEYF; ws.cell(rr, c).font = Font(name='Calibri', size=11, bold=True)
N['dcavg'] = f'$G${rr}'

# ---------------- G. 판정 ----------------
sec('G. 판정 지표 (이수자 평균, 1인당) — B는 형사사법 부분만 반영한 하한')
row('ΔC', '1인당 증분비용(평균)', '={dcavg}', '원/인', 'UC-06 + UC-10(0) + UC-12(0) + ΔM×UC-02(0) → 하한', name='dc', hl=True)
row('B', '1인당 편익(현재가치) = r0×Δr×SCᵣ×D', '={avoid}*{scr}*{D}', '원/인', 'SCᵣ이 형사사법 부분뿐이므로 하한', name='B', hl=True)
row('B/C', '비용편익비율', '={B}/{dc}', '', '', fmt='0.00', name='bc', hl=True)
row('NPV', '순현재가치(1인당)', '={B}-{dc}', '원/인', '', fmt='#,##0;(#,##0)', name='npv')
row('ROI', '투자수익률', '=({B}-{dc})/{dc}', '', '', fmt='0%', name='roi')
row('Δr*', '손익분기 재범감소폭 = ΔC/(r0×SCᵣ×D)', '={dc}/({r0}*{scr}*{D})', '', '이 값보다 실제 Δr이 크면 B/C ≥ 1', fmt='0.0%', name='drstar', hl=True)
row('SCᵣ*', '손익분기 재범 1건당 사회적 비용 = ΔC/(r0×Δr×D)', '={dc}/({avoid}*{D})', '원/건', '3장 SCᵣ가 이 값을 넘으면 B/C ≥ 1', name='scrstar', hl=True)
row('총량', '2025 이수자 전체 증분비용', '={dc}*{comp}', '원', 'PG 비용풀(+UC-10·12 반영분)과 같음', name='totc')
row('총량', '2025 이수자 전체 편익(하한)', '={B}*{comp}', '원', '', name='totb')

# ---------------- H. 민감도 ----------------
sec('H. 민감도 — Δr 시나리오(상대 감소율)별 B/C (다른 값은 위와 동일)')
for lab, v in [('보수', 0.10), ('기본(논문 단순비교)', None), ('중간', 0.25), ('낙관', 0.40)]:
    val = v if v is not None else '={dr}'
    rr0 = row('Δr', f'{lab}', val, '', '', 'assump' if v is not None else 'formula', fmt='0.0%')
    c = ws.cell(rr0, 6, ref(f'={{r0}}*C{rr0}*{{scr}}*{{D}}/{{dc}}')); c.number_format = '0.00'
    ws.cell(rr0, 5, 'B/C →').font = Font(name=F, size=8, color='FF595959')

# ---- 회신 전 대체값(F열)과 자동 연결 산식 ----
def blue(cellref, v, fmt='#,##0', fill=False):
    c = ws[cellref]; c.value = v; c.number_format = fmt
    c.font = Font(name='Calibri', size=11, color=COL['input']);
    if fill: c.fill = YEL
for k, v in [('n1', 439), ('e1', 178), ('n0', 492), ('e0', 247)]:
    blue(N[k].replace('$C$', 'F'), v)
r_n1 = int(N['n1'].split('$')[-1]); ws.cell(r_n1 - 2, 6, '논문값(F)').font = Font(name=F, size=9, bold=True)
u10 = int(N['uc10'].split('$')[-1]); u12 = int(N['uc12'].split('$')[-1])
blue(f'F{u10}', 0, fill=True); blue(f'F{u12}', 0, fill=True); blue(f'G{u12}', 1, fmt='0%', fill=True)
ws.cell(u10 - 1, 6, '수기값(F)').font = Font(name=F, size=9, bold=True)
ws.cell(u12 - 1, 7, '연계비율(G)').font = Font(name=F, size=9, bold=True)
ws[N['uc10'].replace('$', '')].value = ref(f"=IF(AND(ISNUMBER({S7}!H12),ISNUMBER({S7}!H18),N({S7}!H18)>0),{S7}!H12*{{medunit}}/{S7}!H18,F{u10})")
ws[N['uc12'].replace('$', '')].value = ref(f"=IF(AND(ISNUMBER({SX}!H8),ISNUMBER({SX}!H9),N({SX}!H9)>0),{SX}!H8/{SX}!H9*G{u12}/(1+{{rate}}),F{u12})")
for k in ('uc10', 'uc12'):
    c = ws[N[k].replace('$', '')]; c.font = Font(name='Calibri', size=11, color=COL['formula']); c.fill = PatternFill(fill_type=None)
import json
json.dump(N, open('/home/claude/out/calc_names.json', 'w'), ensure_ascii=False, indent=0)

ws.freeze_panes = 'A4'
wb.save(OUT)
print(N)
